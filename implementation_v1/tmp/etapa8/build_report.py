from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).parent / "relatorio_ep1.docx"
doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.2)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.4)
section.right_margin = Cm(2.4)
styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size in [('Title', 16), ('Heading 1', 12)]:
    s = styles[name]
    s.font.name = 'Arial'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor(0,0,0)
    s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(5)
styles['Title'].paragraph_format.space_before = Pt(0)

def p(text=''):
    return doc.add_paragraph(text)
def h(text):
    doc.add_paragraph(text, style='Heading 1')
def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i,x in enumerate(headers): t.rows[0].cells[i].text = str(x)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text = str(x)
    for cell in t.rows[0].cells:
        for para in cell.paragraphs:
            for run in para.runs: run.bold=True
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(2)
                for run in para.runs: run.font.size=Pt(9)
    doc.add_paragraph()

title=doc.add_paragraph(style='Title')
title.add_run('Relatório do EP1 de classificação da clareza das respostas e SIC')
p('ACH2118 — Processamento de Linguagem Natural')
p('Modelo final: BERTimbau Base com ajuste fino para três classes de clareza.')

h('1 Integrantes')
table(['Nome completo', 'Número USP'], [['[NOME COMPLETO DO INTEGRANTE 1]', '[NÚMERO USP]'], ['[NOME COMPLETO DO INTEGRANTE 2]', '[NÚMERO USP]']])
p('Substituir os campos acima pelos integrantes efetivos antes da entrega.')

h('2 Estratégia e modelo final')
p('A tarefa classifica cada resposta textual do e-SIC em c1, c234 ou c5. O modelo final é o BERTimbau Base (neuralmind/bert-base-portuguese-cased) com uma cabeça nova de classificação em três classes. O texto original da coluna resp_text é tokenizado pelo vocabulário do modelo; a sequência usa até 512 tokens, preservando o início quando precisa truncar. A classificação usa os logits da cabeça treinada.')
p('A escolha foi fixada pela maior acurácia média em três folds internos agrupados. A diferença para a referência TF-IDF com regressão logística foi pequena; os dados observados não sustentam uma alegação de superioridade geral do BERTimbau.')

h('3 Pré processamento')
p('Todas as 20.092 linhas e seus rótulos originais foram mantidos, inclusive duplicatas e grupos com rótulos conflitantes. Não houve remoção de stopwords, URLs ou acentos, nem normalização do texto apresentado ao modelo. Valores não textuais em resp_text são convertidos em string. O tokenizer adiciona tokens especiais, trunca após 512 tokens e usa preenchimento dinâmico por lote.')
p('Somente para formar grupos da divisão, cada texto foi convertido para Unicode NFC, casefold e espaços em branco colapsados; o SHA-256 desse texto normalizado identificou o grupo. Nenhum grupo aparece em treino e validação da mesma divisão.')

h('4 Busca de parâmetros')
p('Oito configurações foram comparadas em um subconjunto fixo do desenvolvimento. As duas melhores opções seguiram para a confirmação em três folds internos. A tabela mostra as dimensões da busca e os valores fixos relevantes.')
table(['Parâmetro', 'Valores examinados ou fixados'], [
    ['Representação', 'início; início + fim'],
    ['Comprimento máximo', '256; 512 tokens'],
    ['Taxa de aprendizado', '2 × 10⁻⁵; 3 × 10⁻⁵'],
    ['Otimizador e regularização', 'AdamW; weight decay 0,01'],
    ['Tamanho do lote', 'microbatch 2; acumulação 4; efetivo 8'],
    ['Seleção', 'maior acurácia de validação; macro F1 em empate exato'],
])

h('5 Valores finais')
p('Representação do início com 512 tokens; taxa 2 × 10⁻⁵; AdamW com weight decay 0,01; semente 20260924; lote efetivo 8; 4.523 atualizações do otimizador, aproximadamente 1,80 época no conjunto completo. Não se usa scheduler, warmup, validação durante o ajuste final nem parada antecipada. A configuração completa e a revisão fixada do modelo base constam em configs/bert_stage5_final.json.')

h('6 Divisão e procedimento')
p('Um holdout externo de 3.024 respostas foi separado antes da seleção. As 17.068 respostas de desenvolvimento alimentaram três folds de StratifiedGroupKFold, com grupos disjuntos. O candidato escolhido obteve acurácia média de 45,007% e macro F1 média de 44,690% nesses folds. A referência TF-IDF, C=0,5, atingiu respectivamente 44,705% e 44,453%. A seleção precedeu a única auditoria no holdout.')
p('Após a auditoria, o mesmo protocolo congelado treinou uma nova cabeça de classificação sobre todas as 20.092 respostas por 4.523 passos. Essa execução integral produziu o checkpoint usado para preencher o arquivo de teste. O checkpoint integral contém as linhas antes reservadas ao holdout; por isso a métrica do holdout pertence ao checkpoint de auditoria, treinado somente no desenvolvimento, e não deve ser atribuída ao checkpoint integral.')

h('7 Resultados do modelo escolhido')
p('A tabela apresenta somente o BERTimbau escolhido, avaliado uma vez no holdout externo de 3.024 respostas, antes do treinamento integral.')
table(['Métrica', 'Resultado'], [
    ['Acurácia', '45,966%'], ['Macro F1', '44,523%'], ['Recall c1 (n = 952)', '51,261%'],
    ['Recall c234 (n = 1.036)', '24,131%'], ['Recall c5 (n = 1.036)', '62,934%'],
])
p('O baixo recall de c234 é a principal limitação. Numa comparação pareada adicional solicitada após a auditoria, o TF-IDF obteve 45,437% de acurácia e 45,056% de macro F1 no mesmo holdout. A vantagem de 0,529 ponto percentual do BERTimbau em acurácia corresponde a 16 acertos adicionais; o teste exato de McNemar resultou em p = 0,590. O arquivo test1.xlsx tem 900 respostas sem rótulos, portanto não permite calcular acurácia de teste.')

h('8 Repositório de código')
p('https://github.com/Zennitte/ep1-claridade-esic-bertimbau')
p('O repositório público contém código, parâmetros congelados, versões e instruções. As planilhas, predições e pesos do modelo não são publicados.')

h('9 Reprodução')
for s in [
    '1. Obter train.xlsx e test1.xlsx pelo canal autorizado da disciplina e posicioná-los na raiz do projeto. A planilha de treino esperada tem SHA-256 0e9219233664ac675fc47982bbc822acd3bb15d67fb7bfc26b67bec047e46818.',
    '2. Preparar Python 3.14 e instalar as versões de requirements.lock.txt, com PyTorch ROCm compatível. O ambiente original foi Windows 11, Radeon RX 7600, PyTorch 2.13.0+rocm10.0.0 e Transformers 4.56.2.',
    '3. Executar python src/audit_and_split.py para reconstruir a divisão e verificar os hashes esperados.',
    '4. Executar python src/run_stage6_final.py --run-id reproducao_01 para ajustar o modelo integral em GPU ROCm. A execução salva models/bert_final/ e exige aproximadamente 57 minutos no ambiente original.',
    '5. Executar python src/predict_test_xlsx.py test1.xlsx test1_predito.xlsx --sheet test1 --device cuda. O script preenche clarity na cópia e preserva resp_text.',
]: p(s)
p('A análise do holdout pode ser reproduzida com src/run_stage5f_holdout_audit.py, sempre antes do treino integral, conforme a configuração congelada. A inferência no teste não fornece métricas sem rótulos verdadeiros.')

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.text = 'ACH2118  |  EP1  |  Relatório técnico'
doc.save(OUT)
print(OUT)
