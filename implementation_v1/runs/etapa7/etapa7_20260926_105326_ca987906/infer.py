from __future__ import annotations
import hashlib, json, time
from pathlib import Path
import openpyxl
import torch
from transformers import AutoModelForSequenceClassification, AutoTokenizer, DataCollatorWithPadding

root = Path(__file__).resolve().parents[3]
run_id = Path(__file__).resolve().parent.name
run_dir = root / 'runs' / 'etapa7' / run_id
source = root / 'test1.xlsx'
model_dir = root / 'models' / 'bert_final'

def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024*1024), b''):
            h.update(chunk)
    return h.hexdigest()

start = time.perf_counter()
source_hash = sha256(source)
wb = openpyxl.load_workbook(source, read_only=True, data_only=False)
ws = wb['test1']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
assert headers == ['resp_text', 'clarity'], headers
texts = [row[0] for row in ws.iter_rows(min_row=2, values_only=True)]
assert len(texts) == 900, len(texts)
assert all(isinstance(t, str) and t.strip() for t in texts), 'texto em branco/inválido'
# The target column is read only to verify it contains no test labels; it is never used by the model.
targets = [row[1] for row in ws.iter_rows(min_row=2, values_only=True)]
assert all(v is None for v in targets), 'test labels detected; refusing any evaluation/selection'
wb.close()

device = torch.device('cpu')
tokenizer = AutoTokenizer.from_pretrained(model_dir, local_files_only=True)
model = AutoModelForSequenceClassification.from_pretrained(model_dir, local_files_only=True).to(device)
model.eval()
collator = DataCollatorWithPadding(tokenizer=tokenizer, return_tensors='pt')
id2label = {int(i): label for i, label in model.config.id2label.items()}
assert set(id2label.values()) == {'c1','c234','c5'}
predictions, confidence = [], []
with torch.inference_mode():
    for start_i in range(0, len(texts), 8):
        batch_texts = texts[start_i:start_i+8]
        encoded = tokenizer(batch_texts, truncation=True, max_length=512, padding=False)
        features = [{k: encoded[k][j] for k in encoded} for j in range(len(batch_texts))]
        batch = {k: v.to(device) for k,v in collator(features).items()}
        probs = model(**batch).logits.float().softmax(dim=-1).cpu()
        vals, ids = probs.max(dim=-1)
        predictions.extend(id2label[int(i)] for i in ids.tolist())
        confidence.extend(float(v) for v in vals.tolist())
assert len(predictions) == 900
assert set(predictions) <= set(id2label.values())
from collections import Counter
elapsed = time.perf_counter() - start
artifact_hash = sha256(model_dir / 'model.safetensors')
result = {
  'run_id': run_id,
  'source_file': 'test1.xlsx', 'source_sha256': source_hash,
  'source_sheet': 'test1', 'n_rows': len(texts), 'columns': headers,
  'target_column_initially_empty': True,
  'model_dir': 'models/bert_final', 'model_sha256': artifact_hash,
  'device': str(device), 'batch_size': 8, 'max_length': 512,
  'id2label': id2label,
  'predictions': predictions,
  'counts': dict(sorted(Counter(predictions).items())),
  'mean_max_probability': sum(confidence)/len(confidence),
  'inference_wall_seconds': elapsed,
  'test_labels_used': False,
}
(run_dir / 'inference.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
(run_dir / 'inference.log').write_text(
  f"rows={len(texts)}\ndevice={device}\nbatch_size=8\nmax_length=512\nwall_seconds={elapsed:.3f}\npredicted_counts={dict(Counter(predictions))}\ntest_labels_used=false\nsource_sha256={source_hash}\nmodel_sha256={artifact_hash}\n", encoding='utf-8')
print(json.dumps({k:v for k,v in result.items() if k != 'predictions'}, ensure_ascii=False, indent=2))


