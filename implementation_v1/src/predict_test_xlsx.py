"""Preenche a coluna clarity de uma planilha de teste com o modelo final local."""

from __future__ import annotations

import argparse
from pathlib import Path

import openpyxl
import torch
from transformers import AutoModelForSequenceClassification, AutoTokenizer, DataCollatorWithPadding


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, help="Planilha com resp_text e clarity")
    parser.add_argument("output", type=Path, help="Destino; deve diferir da entrada")
    parser.add_argument("--model", type=Path, default=Path("models/bert_final"))
    parser.add_argument("--sheet", default=None, help="Aba; por padrão, usa a aba ativa")
    parser.add_argument("--batch-size", type=int, default=8)
    parser.add_argument("--device", choices=("auto", "cpu", "cuda"), default="auto")
    args = parser.parse_args()
    if args.batch_size < 1:
        parser.error("--batch-size deve ser positivo")
    if args.input.resolve() == args.output.resolve():
        parser.error("A saída deve diferir da planilha de entrada")
    if args.device == "cuda" and not torch.cuda.is_available():
        parser.error("GPU CUDA/ROCm indisponível")

    book = openpyxl.load_workbook(args.input)
    sheet = book[args.sheet] if args.sheet else book.active
    headers = {str(cell.value).strip(): cell.column for cell in sheet[1] if cell.value is not None}
    if "resp_text" not in headers or "clarity" not in headers:
        parser.error("A primeira linha precisa conter resp_text e clarity")
    text_col, label_col = headers["resp_text"], headers["clarity"]
    rows = [row for row in range(2, sheet.max_row + 1) if sheet.cell(row, text_col).value is not None]
    if any(sheet.cell(row, label_col).value not in (None, "") for row in rows):
        parser.error("A coluna clarity já contém valores; entrada preservada")

    device = torch.device("cuda:0" if args.device == "cuda" or (args.device == "auto" and torch.cuda.is_available()) else "cpu")
    tokenizer = AutoTokenizer.from_pretrained(args.model, local_files_only=True)
    model = AutoModelForSequenceClassification.from_pretrained(args.model, local_files_only=True).to(device)
    model.eval()
    id2label = {int(key): value for key, value in model.config.id2label.items()}
    collator = DataCollatorWithPadding(tokenizer=tokenizer, return_tensors="pt")
    for start in range(0, len(rows), args.batch_size):
        batch_rows = rows[start:start + args.batch_size]
        texts = [str(sheet.cell(row, text_col).value) for row in batch_rows]
        encoded = tokenizer(texts, truncation=True, max_length=512, padding=False)
        features = [{key: encoded[key][i] for key in encoded} for i in range(len(texts))]
        inputs = {key: value.to(device) for key, value in collator(features).items()}
        with torch.inference_mode():
            predictions = model(**inputs).logits.argmax(dim=-1).cpu().tolist()
        for row, pred in zip(batch_rows, predictions):
            sheet.cell(row, label_col).value = id2label[int(pred)]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    print(f"{len(rows)} previsões gravadas em {args.output}")


if __name__ == "__main__":
    main()
