from __future__ import annotations
import csv, hashlib, json, platform, time
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path

import joblib
import numpy as np
import openpyxl
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import Pipeline
from threadpoolctl import threadpool_limits

from audit_and_split import digest_file
from experiment_metrics import classification_metrics

RUN_ID = 'etapa5f_tfidf_holdout_20260926_1115'

def sha256(path: Path) -> str:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda:f.read(1024*1024), b''):
            h.update(block)
    return h.hexdigest()

def main() -> None:
    root=Path(__file__).resolve().parents[1]
    split_path=root/'data'/'splits_grouped_v1.json'
    source_path=root/'train.xlsx'
    config_path=root/'configs'/'baseline.json'
    audit_path=root/'runs'/'etapa5f'/'etapa5f_start512_lr2e5_audit_20260926'/'audit.json'
    bert_csv=root/'runs'/'etapa5f'/'etapa5f_start512_lr2e5_audit_20260926'/'holdout_predictions.csv'
    split=json.loads(split_path.read_text(encoding='utf-8'))
    config=json.loads(config_path.read_text(encoding='utf-8'))
    audit=json.loads(audit_path.read_text(encoding='utf-8'))
    if split['id']!='grouped_v1_seed_20260924' or audit['integrity']['split_sha256']!=digest_file(split_path):
        raise ValueError('A divisão não corresponde à avaliação BERT registrada')
    if audit['integrity']['group_overlap']!=0 or audit['integrity']['index_overlap']!=0:
        raise ValueError('A auditoria BERT registra sobreposição entre desenvolvimento e holdout')
    if digest_file(source_path)!=split['source_sha256'] or audit['integrity']['source_sha256']!=split['source_sha256']:
        raise ValueError('Hash da fonte não confere')
    if config['split_id']!=split['id'] or 0.5 not in config['classifier']['C_values']:
        raise ValueError('Baseline congelado não contém C=0,5')
    dev=np.asarray(split['development'],dtype=np.int64)
    holdout=np.asarray(split['holdout'],dtype=np.int64)
    if len(dev)!=17068 or len(holdout)!=3024 or not set(dev).isdisjoint(holdout):
        raise ValueError('Tamanhos ou disjunção de índices inesperados')
    run_dir=root/'runs'/'etapa5f'/RUN_ID
    model_dir=root/'models'/'baseline'/RUN_ID
    run_dir.mkdir(parents=True,exist_ok=False)
    model_dir.mkdir(parents=True,exist_ok=False)
    started=time.perf_counter()
    wb=openpyxl.load_workbook(source_path,read_only=True,data_only=True)
    ws=wb['train']
    if tuple(c.value for c in next(ws.iter_rows(min_row=1,max_row=1)))!=('resp_text','clarity'):
        raise ValueError('Cabeçalho inesperado')
    rows=list(ws.iter_rows(min_row=2,max_col=2,values_only=True))
    wb.close()
    texts=np.asarray([r[0] if isinstance(r[0],str) else str(r[0]) for r in rows],dtype=object)
    labels=np.asarray([r[1] for r in rows],dtype=object)
    if len(texts)!=20092 or len(labels)!=20092 or set(labels)!=set(config['classes']):
        raise ValueError('Fonte/contagens de classe inesperadas')
    x_train_text,y_train=texts[dev],labels[dev]
    x_hold_text,y_hold=texts[holdout],labels[holdout]
    if set(y_train)!=set(config['classes']) or set(y_hold)!=set(config['classes']):
        raise ValueError('Alguma classe ausente da partição')

    bert_by_index={}
    with bert_csv.open(encoding='utf-8-sig',newline='') as f:
        for row in csv.DictReader(f):
            idx=int(row['row_index'])
            if idx in bert_by_index: raise ValueError(f'Índice BERT duplicado: {idx}')
            bert_by_index[idx]=(row['true_label'],row['predicted_label'])
    if set(bert_by_index)!=set(holdout.tolist()):
        raise ValueError('Predições BERT não correspondem aos mesmos índices de holdout')
    bert_true=np.asarray([bert_by_index[int(i)][0] for i in holdout],dtype=object)
    bert_pred=np.asarray([bert_by_index[int(i)][1] for i in holdout],dtype=object)
    if not np.array_equal(bert_true,y_hold): raise ValueError('Rótulos BERT divergem do holdout da divisão')
    bert_metrics=classification_metrics(y_hold,bert_pred,config['classes'])
    recorded=audit['holdout_metrics']
    if abs(bert_metrics['accuracy']-recorded['accuracy'])>1e-12:
        raise ValueError('Métrica BERT recomputada diverge do audit.json')

    vector_config=dict(config['vectorizer'])
    vector_config['ngram_range']=tuple(vector_config['ngram_range'])
    vectorizer=TfidfVectorizer(**vector_config)
    with threadpool_limits(limits=config['cpu_threads']):
        x_train=vectorizer.fit_transform(x_train_text)
        x_hold=vectorizer.transform(x_hold_text)
        clf=LogisticRegression(C=0.5,solver=config['classifier']['solver'],max_iter=config['classifier']['max_iter'],random_state=config['seed'])
        clf.fit(x_train,y_train)
        pred=clf.predict(x_hold)
        probs=clf.predict_proba(x_hold)
    baseline_metrics=classification_metrics(y_hold,pred,config['classes'])
    pipeline=Pipeline([('tfidf',vectorizer),('logreg',clf)])
    model_path=model_dir/'tfidf_logreg_C_0_5.joblib'
    joblib.dump(pipeline,model_path,compress=3)
    comparison={
      'stage':'5f_addendum','run_id':RUN_ID,'created_utc':datetime.now(timezone.utc).isoformat(),
      'source_sha256':digest_file(source_path),'split_id':split['id'],'split_sha256':digest_file(split_path),
      'baseline_config_sha256':digest_file(config_path),'bert_run_id':audit['run_id'],
      'development_rows':len(dev),'holdout_rows':len(holdout),'same_holdout_indices_as_bert':True,
      'same_training_indices_as_bert':True,'holdout_groups_overlap_development':False,
      'selected_C':0.5,'selected_before_holdout_from_grouped_cv':True,
      'vectorizer':config['vectorizer'],'classifier':{'solver':config['classifier']['solver'],'max_iter':config['classifier']['max_iter'],'random_state':config['seed']},
      'vocabulary_size':len(vectorizer.vocabulary_),'train_nonzero':int(x_train.nnz),'holdout_nonzero':int(x_hold.nnz),
      'classes':config['classes'],'baseline_metrics':baseline_metrics,'bert_metrics_recomputed':bert_metrics,
      'difference_bert_minus_tfidf_percentage_points':{
        'accuracy':100*(bert_metrics['accuracy']-baseline_metrics['accuracy']),
        'macro_f1':100*(bert_metrics['macro_f1']-baseline_metrics['macro_f1'])},
      'model_path':str(model_path.relative_to(root)),'model_sha256':sha256(model_path),
      'rows':len(holdout),'versions':{n:version(n) for n in ('scikit-learn','numpy','openpyxl','joblib')},
      'platform':platform.platform(),'cpu_threads':config['cpu_threads'],'gpu_used':False,
      'wall_seconds':round(time.perf_counter()-started,3),
      'holdout_used_for_tuning_or_selection':False,
    }
    (run_dir/'comparison.json').write_text(json.dumps(comparison,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with (run_dir/'predictions.csv').open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f); w.writerow(['row_index','true_label','tfidf_prediction','bert_prediction','tfidf_confidence'])
        for i,true,pred_i,bert_i,p in zip(holdout,y_hold,pred,bert_pred,probs.max(axis=1)):
            w.writerow([int(i),true,pred_i,bert_i,f'{p:.9f}'])
    print(json.dumps(comparison,ensure_ascii=False,indent=2))

if __name__=='__main__': main()


