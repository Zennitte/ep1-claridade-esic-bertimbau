"""Verifica as entradas e a capacidade de treino usada no EP1."""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import sys
import time
from datetime import datetime, timezone
from importlib.metadata import version
from pathlib import Path


MODEL_ID = "neuralmind/bert-base-portuguese-cased"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def run(root: Path) -> dict:
    import openpyxl
    import torch
    from pypdf import PdfReader
    from transformers import AutoTokenizer

    inputs = {}
    for filename in ("train.xlsx", "ep1-enunciado.pdf"):
        path = root / filename
        inputs[filename] = {"size_bytes": path.stat().st_size, "sha256": sha256(path)}

    workbook = openpyxl.load_workbook(root / "train.xlsx", read_only=True, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    inputs["train.xlsx"].update(
        {"sheet": sheet.title, "rows_including_header": sheet.max_row, "columns": headers}
    )
    workbook.close()

    pdf = PdfReader(root / "ep1-enunciado.pdf")
    page_texts = [(page.extract_text() or "") for page in pdf.pages]
    inputs["ep1-enunciado.pdf"].update(
        {"pages": len(pdf.pages), "extractable_characters": sum(map(len, page_texts))}
    )

    tokenizer = AutoTokenizer.from_pretrained(MODEL_ID)
    token_ids = tokenizer("Esta é uma resposta de verificação.")["input_ids"]

    gpu = {
        "available": torch.cuda.is_available(),
        "torch_hip": torch.version.hip,
        "torch_cuda": torch.version.cuda,
    }
    if not gpu["available"]:
        raise RuntimeError("PyTorch não detectou uma GPU compatível")

    device = torch.device("cuda:0")
    gpu["name"] = torch.cuda.get_device_name(device)
    props = torch.cuda.get_device_properties(device)
    gpu["total_memory_bytes"] = props.total_memory
    torch.manual_seed(20260924)
    torch.cuda.reset_peak_memory_stats(device)
    model = torch.nn.Sequential(
        torch.nn.Embedding(tokenizer.vocab_size, 32),
        torch.nn.Flatten(),
        torch.nn.Linear(32 * 16, 3),
    ).to(device)
    optimizer = torch.optim.AdamW(model.parameters(), lr=1e-3)
    batch = torch.randint(0, tokenizer.vocab_size, (8, 16), device=device)
    labels = torch.randint(0, 3, (8,), device=device)
    before = model[-1].weight.detach().clone()
    torch.cuda.synchronize(device)
    start = time.perf_counter()
    optimizer.zero_grad(set_to_none=True)
    logits = model(batch)
    loss = torch.nn.functional.cross_entropy(logits, labels)
    loss.backward()
    optimizer.step()
    torch.cuda.synchronize(device)
    gpu.update(
        {
            "training_step_seconds": round(time.perf_counter() - start, 6),
            "peak_allocated_bytes": torch.cuda.max_memory_allocated(device),
            "peak_reserved_bytes": torch.cuda.max_memory_reserved(device),
            "loss": float(loss.item()),
            "weights_changed": bool(torch.any(before != model[-1].weight).item()),
        }
    )
    if not gpu["weights_changed"]:
        raise RuntimeError("O passo de treino não atualizou os pesos")

    return {
        "timestamp_utc": datetime.now(timezone.utc).isoformat(),
        "python": sys.version,
        "platform": platform.platform(),
        "versions": {
            package: version(package)
            for package in ("torch", "transformers", "tokenizers", "openpyxl", "pypdf", "scikit-learn")
        },
        "inputs": inputs,
        "tokenizer": {"model_id": MODEL_ID, "sample_token_count": len(token_ids)},
        "gpu": gpu,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path("runs/etapa0/verification.json"))
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[1]
    result = run(root)
    output = args.output if args.output.is_absolute() else root / args.output
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
