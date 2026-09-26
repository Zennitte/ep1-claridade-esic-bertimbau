"""Treina o modelo de auditoria em desenvolvimento e consulta o holdout uma vez."""

from __future__ import annotations

import argparse
import csv
import json
import random
import time
from datetime import datetime, timezone
from hashlib import sha256
from importlib.metadata import version
from pathlib import Path

import numpy as np
import openpyxl
import torch
from torch.utils.data import DataLoader
from transformers import AutoModelForSequenceClassification, AutoTokenizer, DataCollatorWithPadding

from audit_and_split import digest_file
from bert_stage3 import EncodedDataset
from experiment_metrics import classification_metrics

EXPECTED_SPLIT_SHA256 = "aa9c99a79ccccc3655399761c53098f5fd3d770c9aa6c8b68be29c49eba3586b"
EXPECTED_ROW_MANIFEST_SHA256 = "e07917a05f89f9940d6194d118fed7ebce9b9342acfc956eb38cadfd4dc2a2c4"
EXPECTED_FROZEN_CONFIG_SHA256 = "4377c55014be7f6f3de1b89b9dfb9c64785561b3df1a4b96b4eeef7e9a8b0953"


def write_json(path: Path, obj: dict) -> None:
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_rows(root: Path, source: str, indices: list[int], want_labels: bool) -> tuple[list[str], list[str] | None]:
    """Acessa somente as linhas solicitadas; índice zero corresponde à linha Excel 2."""
    # Non-read-only mode makes indexed access linear overall instead of repeatedly
    # rescanning the worksheet's XML for each of the 17k development indices.
    workbook = openpyxl.load_workbook(root / source, read_only=False, data_only=True)
    sheet = workbook["train"]
    texts: list[str] = []
    labels: list[str] | None = [] if want_labels else None
    for index in indices:
        text = sheet.cell(row=index + 2, column=1).value
        texts.append(text if isinstance(text, str) else str(text))
        if labels is not None:
            labels.append(sheet.cell(row=index + 2, column=2).value)
    workbook.close()
    return texts, labels


def verify(root: Path, config: dict, split: dict, split_path: Path) -> dict:
    if digest_file(root / config["source"]["path"]) != config["source"]["sha256"]:
        raise ValueError("Hash de train.xlsx diverge da especificação congelada")
    if split["source_sha256"] != config["source"]["sha256"]:
        raise ValueError("Hash da fonte na divisão diverge da especificação congelada")
    if digest_file(split_path) != EXPECTED_SPLIT_SHA256:
        raise ValueError("Hash da divisão diverge da especificação congelada")
    dev, holdout = split["development"], split["holdout"]
    if len(dev) != 17068 or len(holdout) != 3024:
        raise ValueError("Contagens de desenvolvimento/holdout divergem do protocolo 5f")
    if set(dev) & set(holdout) or len(set(dev)) != len(dev) or len(set(holdout)) != len(holdout):
        raise ValueError("Índices duplicados ou sobreposição entre desenvolvimento e holdout")
    if set(dev) | set(holdout) != set(range(config["source"]["rows"])):
        raise ValueError("A divisão não cobre exatamente as linhas de train.xlsx")

    manifest_path = root / "data" / "row_manifest.csv"
    if digest_file(manifest_path) != EXPECTED_ROW_MANIFEST_SHA256:
        raise ValueError("Hash de row_manifest.csv diverge")
    groups: dict[int, str] = {}
    dev_set, holdout_set = set(dev), set(holdout)
    with manifest_path.open(encoding="utf-8", newline="") as f:
        next(f)  # header
        for line in f:
            fields = line.rstrip("\r\n").split(",", 3)
            idx = int(fields[0])
            if idx in dev_set or idx in holdout_set:
                groups[idx] = fields[2]
    if len(groups) != len(dev) + len(holdout):
        raise ValueError("Manifesto de grupos não cobre a divisão")
    if {groups[i] for i in dev} & {groups[i] for i in holdout}:
        raise ValueError("Grupo normalizado aparece nos dois lados da divisão")

    dev_labels = read_rows(root, config["source"]["path"], dev, True)[1]
    label_names = [config["model"]["id_to_label"][str(i)] for i in range(config["model"]["num_labels"])]
    counts = {name: dev_labels.count(name) for name in label_names}
    expected_dev_counts = {"c1": 5395, "c234": 5817, "c5": 5856}
    if counts != expected_dev_counts:
        raise ValueError(f"Contagens de desenvolvimento divergentes: {counts}")
    return {
        "source_sha256": digest_file(root / config["source"]["path"]),
        "split_sha256": digest_file(split_path),
        "row_manifest_sha256": digest_file(manifest_path),
        "development_rows": len(dev),
        "development_class_counts": counts,
        "holdout_rows": len(holdout),
        "index_overlap": 0,
        "group_overlap": 0,
        "holdout_labels_requested_during_preflight": False,
    }


def metrics_with_support(y_true: list[str], y_pred: list[str], classes: list[str]) -> dict:
    result = classification_metrics(y_true, y_pred, classes)
    result["support_by_class"] = {label: y_true.count(label) for label in classes}
    return result


def run(root: Path, run_id: str) -> dict:
    config_path = root / "configs" / "bert_stage5_final.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    split_path = root / "data" / "splits_grouped_v1.json"
    split = json.loads(split_path.read_text(encoding="utf-8"))
    if split["id"] != "grouped_v1_seed_20260924":
        raise ValueError("ID da divisão inesperado")
    integrity = verify(root, config, split, split_path)
    if digest_file(config_path) != EXPECTED_FROZEN_CONFIG_SHA256:
        raise ValueError("Especificação 5e foi modificada após o congelamento")

    run_dir = root / "runs" / "etapa5f" / run_id
    checkpoint_dir = root / "models" / "bert_stage5f" / run_id / "final"
    if run_dir.exists() or checkpoint_dir.parent.exists():
        raise FileExistsError("Run/checkpoint já existe; a auditoria não pode ser repetida")
    if not torch.cuda.is_available():
        raise RuntimeError("GPU ROCm indisponível")
    device = torch.device("cuda:0")
    gpu = torch.cuda.get_device_properties(device)

    train_indices = split["development"]
    holdout_indices = split["holdout"]
    labels = [config["model"]["id_to_label"][str(i)] for i in range(config["model"]["num_labels"])]
    label2id = {name: index for index, name in enumerate(labels)}
    id2label = {index: name for name, index in label2id.items()}
    training = config["training"]
    target_steps = round(training["target_epoch_equivalents"] * (len(train_indices) // training["effective_batch_size"]))
    steps_per_epoch = len(train_indices) // training["effective_batch_size"]
    if target_steps != 3842 or steps_per_epoch != 2133:
        raise ValueError(f"Duração calculada inesperada: {target_steps} steps, {steps_per_epoch}/época")

    run_dir.mkdir(parents=True, exist_ok=False)
    checkpoint_dir.mkdir(parents=True, exist_ok=False)
    result = {
        "stage": "5f",
        "run_id": run_id,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "frozen_config_sha256": digest_file(config_path),
        "integrity": integrity,
        "model": config["model"],
        "training": {
            **training,
            "development_steps_per_epoch": steps_per_epoch,
            "target_optimizer_steps_5f": target_steps,
            "actual_epoch_equivalents": target_steps / steps_per_epoch,
        },
        "holdout_evaluation_started": False,
        "holdout_evaluated": False,
        "gpu": {"name": gpu.name, "total_memory_bytes": gpu.total_memory},
        "versions": {n: version(n) for n in ("torch", "transformers", "tokenizers", "openpyxl")},
    }
    write_json(run_dir / "audit.json", result)
    write_json(run_dir / "config.json", config)

    torch.manual_seed(training["seed"])
    np.random.seed(training["seed"])
    random.seed(training["seed"])
    torch.cuda.empty_cache()
    torch.cuda.reset_peak_memory_stats(device)

    dev_texts, dev_labels = read_rows(root, config["source"]["path"], train_indices, True)
    if set(dev_labels) != set(labels):
        raise ValueError("Rótulos de desenvolvimento diferem do mapeamento congelado")
    tokenizer = AutoTokenizer.from_pretrained(
        config["model"]["tokenizer_id"],
        revision=config["model"]["tokenizer_revision"], local_files_only=True,
    )
    enc = tokenizer(dev_texts, truncation=True, max_length=config["model"]["max_length_including_special_tokens"], padding=False)
    dataset = EncodedDataset(enc, [label2id[x] for x in dev_labels])
    train_loader = DataLoader(
        dataset, batch_size=training["micro_batch_size"], shuffle=True,
        generator=torch.Generator().manual_seed(training["shuffle_generator_seed"]),
        num_workers=training["num_workers"],
        collate_fn=DataCollatorWithPadding(tokenizer=tokenizer, return_tensors="pt"),
    )
    model = AutoModelForSequenceClassification.from_pretrained(
        config["model"]["pretrained_model_id"],
        revision=config["model"]["pretrained_model_revision"],
        num_labels=len(labels), label2id=label2id, id2label=id2label,
        local_files_only=True,
    ).to(device)
    optimizer = torch.optim.AdamW(
        model.parameters(), lr=training["learning_rate"],
        betas=tuple(training["adam_betas"]), eps=training["adam_epsilon"],
        weight_decay=training["weight_decay"],
    )
    accumulation = training["gradient_accumulation_steps"]
    max_micro_per_epoch = steps_per_epoch * accumulation
    started = time.perf_counter()
    step = 0
    log_path = run_dir / "training.jsonl"
    with log_path.open("w", encoding="utf-8") as log:
        while step < target_steps:
            model.train()
            optimizer.zero_grad(set_to_none=True)
            for micro_index, batch in enumerate(train_loader):
                if micro_index >= max_micro_per_epoch:
                    break
                batch = {k: v.to(device) for k, v in batch.items()}
                output = model(**batch)
                (output.loss / accumulation).backward()
                if (micro_index + 1) % accumulation == 0:
                    torch.nn.utils.clip_grad_norm_(model.parameters(), training["gradient_clip_norm"])
                    optimizer.step()
                    optimizer.zero_grad(set_to_none=True)
                    step += 1
                    if step % 100 == 0 or step == target_steps:
                        torch.cuda.synchronize(device)
                        entry = {
                            "step": step,
                            "epoch_equivalent": step / steps_per_epoch,
                            "last_microbatch_loss": float(output.loss.item()),
                            "elapsed_seconds": round(time.perf_counter() - started, 2),
                        }
                        log.write(json.dumps(entry) + "\n")
                        log.flush()
                        print(json.dumps(entry), flush=True)
                    if step >= target_steps:
                        break
    torch.cuda.synchronize(device)
    result["training_completed_utc"] = datetime.now(timezone.utc).isoformat()
    result["train_steps_completed"] = step
    result["train_seconds"] = round(time.perf_counter() - started, 3)
    result["peak_gpu_memory_bytes"] = torch.cuda.max_memory_allocated(device)
    model.save_pretrained(checkpoint_dir, safe_serialization=True)
    tokenizer.save_pretrained(checkpoint_dir)
    result["checkpoint_sha256"] = {
        p.name: digest_file(p) for p in sorted(checkpoint_dir.iterdir()) if p.is_file()
    }
    write_json(run_dir / "audit.json", result)

    # The marker is persisted before loading holdout text or labels. Any interrupted
    # inference remains non-repeatable and requires an explicit audit review.
    result["holdout_evaluation_started"] = True
    result["holdout_evaluation_started_utc"] = datetime.now(timezone.utc).isoformat()
    write_json(run_dir / "audit.json", result)
    holdout_texts, holdout_labels = read_rows(root, config["source"]["path"], holdout_indices, True)
    result["integrity"]["holdout_labels_requested_for_audit"] = True
    hold_enc = tokenizer(holdout_texts, truncation=True, max_length=config["model"]["max_length_including_special_tokens"], padding=False)
    hold_dataset = EncodedDataset(hold_enc, [label2id[x] for x in holdout_labels])
    hold_loader = DataLoader(
        hold_dataset, batch_size=training["micro_batch_size"] * 2, shuffle=False,
        num_workers=0, collate_fn=DataCollatorWithPadding(tokenizer=tokenizer, return_tensors="pt"),
    )
    model.eval()
    y_true: list[str] = []
    y_pred: list[str] = []
    confidences: list[float] = []
    probabilities: list[list[float]] = []
    losses: list[float] = []
    audit_started = time.perf_counter()
    with torch.inference_mode():
        for batch in hold_loader:
            batch = {k: v.to(device) for k, v in batch.items()}
            output = model(**batch)
            losses.append(float(output.loss.item()) * len(batch["labels"]))
            probs = output.logits.softmax(dim=-1)
            pred = probs.argmax(dim=-1)
            y_true.extend(labels[int(i)] for i in batch["labels"].tolist())
            y_pred.extend(labels[int(i)] for i in pred.tolist())
            confidences.extend(float(x) for x in probs.max(dim=-1).values.tolist())
            probabilities.extend([[float(v) for v in row] for row in probs.tolist()])
    torch.cuda.synchronize(device)
    result["holdout_evaluated"] = True
    result["holdout_evaluation_completed_utc"] = datetime.now(timezone.utc).isoformat()
    result["holdout_inference_seconds"] = round(time.perf_counter() - audit_started, 3)
    result["holdout_metrics"] = metrics_with_support(y_true, y_pred, labels)
    result["holdout_loss"] = sum(losses) / len(y_true)
    pred_path = run_dir / "holdout_predictions.csv"
    with pred_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["row_index", "true_label", "predicted_label", "confidence", *[f"p_{x}" for x in labels]])
        for row, true, pred, conf, probs in zip(holdout_indices, y_true, y_pred, confidences, probabilities):
            writer.writerow([row, true, pred, f"{conf:.8f}", *[f"{p:.8f}" for p in probs]])
    result["holdout_predictions_sha256"] = digest_file(pred_path)
    write_json(run_dir / "metrics.json", result["holdout_metrics"] | {"loss": result["holdout_loss"], "rows": len(y_true)})
    write_json(run_dir / "audit.json", result)
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--run-id")
    parser.add_argument("--preflight", action="store_true")
    args = parser.parse_args()
    config = json.loads((args.root.resolve() / "configs" / "bert_stage5_final.json").read_text(encoding="utf-8"))
    split_path = args.root.resolve() / "data" / "splits_grouped_v1.json"
    split = json.loads(split_path.read_text(encoding="utf-8"))
    if args.preflight:
        if digest_file(args.root.resolve() / "configs" / "bert_stage5_final.json") != EXPECTED_FROZEN_CONFIG_SHA256:
            raise ValueError("Especificação 5e foi modificada após o congelamento")
        print(json.dumps(verify(args.root.resolve(), config, split, split_path), ensure_ascii=False, indent=2))
        return
    if not args.run_id:
        parser.error("--run-id é obrigatório fora do preflight")
    result = run(args.root.resolve(), args.run_id)
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
