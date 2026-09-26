"""Audita train.xlsx e congela divisões agrupadas para o EP1."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import time
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
from sklearn.model_selection import GroupShuffleSplit, StratifiedGroupKFold
from transformers import AutoTokenizer


SEED = 20260924
LABELS = ("c1", "c234", "c5")
SOURCE = "train.xlsx"
EXPECTED_SHA256 = "0e9219233664ac675fc47982bbc822acd3bb15d67fb7bfc26b67bec047e46818"


def digest_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def group_text(text: str) -> str:
    """Somente Unicode NFC, caixa e sequências de espaços; sem alterar o treino."""
    return " ".join(unicodedata.normalize("NFC", text).casefold().split())


def group_id(normalized_text: str) -> str:
    return hashlib.sha256(normalized_text.encode("utf-8")).hexdigest()


def distribution(labels: np.ndarray, indices: np.ndarray) -> dict:
    counts = Counter(labels[indices].tolist())
    total = len(indices)
    return {
        "rows": total,
        "counts": {label: counts[label] for label in LABELS},
        "shares": {label: round(counts[label] / total, 6) for label in LABELS},
    }


def quantiles(values: list[int]) -> dict:
    return {
        "min": min(values),
        "p10": float(np.percentile(values, 10)),
        "median": float(np.median(values)),
        "p90": float(np.percentile(values, 90)),
        "p99": float(np.percentile(values, 99)),
        "max": max(values),
    }


def safe_excerpt(text: str) -> str:
    text = "".join(" " if unicodedata.category(char) == "Cc" else char for char in text)
    text = re.sub(r"https?://\S+", "[URL]", text)
    text = re.sub(r"[\w.+-]+@[\w.-]+", "[EMAIL]", text)
    text = re.sub(r"\b\d{3}[.\- ]?\d{3}[.\- ]?\d{3}[- ]?\d{2}\b", "[ID]", text)
    text = re.sub(r"\b\d{7,}\b", "[NUMERO]", text)
    return " ".join(text.split())[:180]


def inspect_examples(texts: list[str], labels: np.ndarray, norm: list[str]) -> list[dict]:
    lengths = np.array([len(text.split()) for text in texts])
    choices = [("mais_curto", int(np.argmin(lengths))), ("mais_longo", int(np.argmax(lengths)))]
    for label in LABELS:
        candidates = np.flatnonzero(labels == label)
        typical = candidates[np.argmin(abs(lengths[candidates] - np.median(lengths[candidates])))]
        choices.append((f"mediano_{label}", int(typical)))
    by_group = defaultdict(list)
    for index, key in enumerate(norm):
        by_group[key].append(index)
    for key, indices in by_group.items():
        if len({labels[i] for i in indices}) > 1:
            choices.append(("conflito", indices[0]))
            break
    return [
        {
            "tipo": category,
            "row_index": index,
            "excel_row": index + 2,
            "label": str(labels[index]),
            "word_count": int(lengths[index]),
            "excerpt_masked": safe_excerpt(texts[index]),
        }
        for category, index in choices
    ]


def run(root: Path, holdout_candidates: int) -> dict:
    started = time.perf_counter()
    source = root / SOURCE
    source_hash = digest_file(source)
    if source_hash != EXPECTED_SHA256:
        raise ValueError("O hash de train.xlsx mudou; rever o protocolo antes de dividir")

    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    if workbook.sheetnames != ["train"]:
        raise ValueError(f"Abas inesperadas: {workbook.sheetnames}")
    sheet = workbook["train"]
    rows = sheet.iter_rows(values_only=True)
    header = next(rows)
    if header != ("resp_text", "clarity"):
        raise ValueError(f"Cabeçalho inesperado: {header}")

    texts: list[str] = []
    raw_labels: list[str] = []
    nulls = Counter()
    types = {"resp_text": Counter(), "clarity": Counter()}
    for row_number, row in enumerate(rows, start=2):
        if len(row) != 2:
            raise ValueError(f"Linha {row_number} tem {len(row)} colunas")
        text, label = row
        types["resp_text"][type(text).__name__] += 1
        types["clarity"][type(label).__name__] += 1
        if text is None or (isinstance(text, str) and not text.strip()):
            nulls["empty_text"] += 1
        if label is None or not isinstance(label, str) or not label.strip():
            nulls["empty_or_invalid_label"] += 1
        # Uma célula numérica é mantida como linha de treino, com conversão explícita.
        texts.append(text if isinstance(text, str) else str(text) if text is not None else "")
        raw_labels.append(label if isinstance(label, str) else "")
    workbook.close()
    labels = np.array(raw_labels)
    if nulls or set(raw_labels) != set(LABELS):
        raise ValueError(f"Dados incompatíveis: nulos={dict(nulls)}, classes={sorted(set(raw_labels))}")

    normalized = [group_text(text) for text in texts]
    groups = np.array([group_id(value) for value in normalized])
    assert len(set(groups)) == len(set(normalized)), "Colisão de hash de grupo"
    n = len(texts)
    all_indices = np.arange(n)
    exact = Counter(texts)
    exact_rows = defaultdict(list)
    group_rows = defaultdict(list)
    for index, key in enumerate(groups):
        exact_rows[texts[index]].append(index)
        group_rows[key].append(index)
    conflicting = {
        key: indices for key, indices in group_rows.items() if len({labels[i] for i in indices}) > 1
    }
    words = [len(text.split()) for text in texts]
    chars = [len(text) for text in texts]
    tokenizer = AutoTokenizer.from_pretrained("neuralmind/bert-base-portuguese-cased", local_files_only=True)
    tokens = []
    for start in range(0, len(texts), 256):
        encoded = tokenizer(texts[start : start + 256], add_special_tokens=True, truncation=False)
        tokens.extend(len(ids) for ids in encoded["input_ids"])
    quality = {
        "blank_after_strip": sum(not text.strip() for text in texts),
        "leading_or_trailing_whitespace": sum(text != text.strip() for text in texts),
        "repeated_internal_whitespace": sum(bool(re.search(r"\s{2,}", text)) for text in texts),
        "contains_html_tag": sum(bool(re.search(r"<[^>]+>", text)) for text in texts),
        "contains_url": sum(bool(re.search(r"https?://|www\.", text, re.I)) for text in texts),
        "contains_email": sum(bool(re.search(r"[\w.+-]+@[\w.-]+", text)) for text in texts),
        "contains_control_character": sum(any(unicodedata.category(c) == "Cc" and c not in "\t\n\r" for c in text) for text in texts),
        "at_most_5_words": sum(value <= 5 for value in words),
        "over_512_words": sum(value > 512 for value in words),
    }

    # A escolha usa apenas tamanho e proporções dos rótulos, sem olhar desempenho.
    global_shares = np.array([np.mean(labels == label) for label in LABELS])
    splitter = GroupShuffleSplit(n_splits=holdout_candidates, test_size=0.15, random_state=SEED)
    candidates = []
    for candidate_id, (development, holdout) in enumerate(splitter.split(all_indices, labels, groups)):
        shares = np.array([np.mean(labels[holdout] == label) for label in LABELS])
        size_error = abs(len(holdout) / n - 0.15)
        class_error = float(np.max(abs(shares - global_shares)))
        score = size_error + class_error
        candidates.append((score, size_error, class_error, candidate_id, development, holdout))
    score, size_error, class_error, selected_id, development, holdout = min(candidates, key=lambda item: item[:4])
    if size_error > 0.01 or class_error > 0.02:
        raise ValueError(f"Holdout desbalanceado: size_error={size_error}, class_error={class_error}")

    folds = []
    fold_splitter = StratifiedGroupKFold(n_splits=3, shuffle=True, random_state=SEED)
    for fold_number, (train_rel, valid_rel) in enumerate(
        fold_splitter.split(development, labels[development], groups[development]), start=1
    ):
        train = development[train_rel]
        valid = development[valid_rel]
        folds.append({"fold": fold_number, "train": train.tolist(), "validation": valid.tolist()})

    holdout_groups = set(groups[holdout])
    assert not holdout_groups.intersection(groups[development])
    assert set(development).isdisjoint(holdout)
    assert set(development).union(holdout) == set(all_indices)
    validation_sets = []
    for fold in folds:
        train = np.array(fold["train"])
        valid = np.array(fold["validation"])
        assert not set(groups[train]).intersection(groups[valid])
        assert set(train).isdisjoint(valid)
        assert set(train).union(valid) == set(development)
        validation_sets.append(set(valid))
    assert set.union(*validation_sets) == set(development)
    assert sum(map(len, validation_sets)) == len(development)

    reports = root / "reports"
    data = root / "data"
    reports.mkdir(exist_ok=True)
    data.mkdir(exist_ok=True)
    with (data / "row_manifest.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(("row_index", "excel_row", "group_sha256", "clarity"))
        writer.writerows((i, i + 2, groups[i], labels[i]) for i in all_indices)

    split = {
        "id": "grouped_v1_seed_20260924",
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "source": SOURCE,
        "source_sha256": source_hash,
        "row_indexing": "zero_based_data_rows; Excel row = row_index + 2",
        "normalization": "Unicode NFC, casefold, split/join all whitespace",
        "seed": SEED,
        "holdout_method": "GroupShuffleSplit(test_size=0.15, n_splits=256); lowest sum of absolute size error and maximum class-share error",
        "holdout_candidate_count": holdout_candidates,
        "selected_candidate": selected_id,
        "holdout_used_for_model_selection": False,
        "development": development.tolist(),
        "holdout": holdout.tolist(),
        "fold_method": "StratifiedGroupKFold(n_splits=3, shuffle=True)",
        "folds": folds,
    }
    split_path = data / "splits_grouped_v1.json"
    split_path.write_text(json.dumps(split, ensure_ascii=False) + "\n", encoding="utf-8")

    summary = {
        "source": SOURCE,
        "source_sha256": source_hash,
        "sheet": "train",
        "headers": list(header),
        "rows": n,
        "types": {column: dict(counts) for column, counts in types.items()},
        "nulls": dict(nulls),
        "classes": distribution(labels, all_indices),
        "exact_distinct_texts": len(exact),
        "exact_repeated_groups": sum(count > 1 for count in exact.values()),
        "exact_repeated_rows": sum(count for count in exact.values() if count > 1),
        "normalized_distinct_texts": len(group_rows),
        "normalized_repeated_groups": sum(len(indices) > 1 for indices in group_rows.values()),
        "normalized_repeated_rows": sum(len(indices) for indices in group_rows.values() if len(indices) > 1),
        "conflicting_groups": len(conflicting),
        "conflicting_rows": sum(map(len, conflicting.values())),
        "exact_conflicting_groups": sum(
            len({labels[i] for i in indices}) > 1 for indices in exact_rows.values()
        ),
        "same_normalized_text_majority_ceiling_rows": sum(
            max(Counter(labels[indices]).values()) for indices in group_rows.values()
        ),
        "largest_group_rows": max(map(len, group_rows.values())),
        "word_count": quantiles(words),
        "character_count": quantiles(chars),
        "bertimbau_token_count_with_special_tokens": quantiles(tokens),
        "over_256_bertimbau_tokens": sum(value > 256 for value in tokens),
        "over_512_bertimbau_tokens": sum(value > 512 for value in tokens),
        "quality_flags": quality,
        "holdout": distribution(labels, holdout),
        "development": distribution(labels, development),
        "holdout_size_error": float(size_error),
        "holdout_max_class_share_error": float(class_error),
        "holdout_candidate": selected_id,
        "folds": [
            {
                "fold": fold["fold"],
                "train": distribution(labels, np.array(fold["train"])),
                "validation": distribution(labels, np.array(fold["validation"])),
            }
            for fold in folds
        ],
        "leakage_checks": {
            "holdout_group_overlap": 0,
            "fold_group_overlap": [0] * len(folds),
            "all_rows_assigned_once_to_holdout_or_development": True,
            "each_development_row_validated_once": True,
        },
        "representative_examples": inspect_examples(texts, labels, normalized),
        "elapsed_seconds": round(time.perf_counter() - started, 3),
    }
    (reports / "etapa1_audit.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--holdout-candidates", type=int, default=256)
    args = parser.parse_args()
    if args.holdout_candidates < 1:
        parser.error("--holdout-candidates deve ser positivo")
    summary = run(Path(__file__).resolve().parents[1], args.holdout_candidates)
    print(json.dumps(summary, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
